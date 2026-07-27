"""中文 PowerPoint 导出器 —— 基于 python-pptx"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from typing import List
from src.schemas.lesson import 课件页面, 课程规范


class 中文PPT导出器:
    """中文 PowerPoint 导出器"""

    # 中国传统配色 —— 千里江山图风格
    主题色_青 = RGBColor(0x2E, 0x75, 0xB6)    # 主色：青蓝色
    主题色_绿 = RGBColor(0x5B, 0x8C, 0x5A)    # 辅助色：苍绿色
    主题色_金 = RGBColor(0xE8, 0xA8, 0x3E)    # 强调色：赭金色
    主题色_灰 = RGBColor(0x4A, 0x4A, 0x4A)    # 正文深灰
    主题色_浅灰 = RGBColor(0xF5, 0xF5, 0xF0)  # 背景浅米
    主题色_白 = RGBColor(0xFF, 0xFF, 0xFF)

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)  # 16:9 宽屏
        self.prs.slide_height = Inches(7.5)

    def _设置幻灯片背景(self, slide, 颜色):
        """设置幻灯片纯色背景"""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = 颜色

    def _添加文本框(self, slide, left, top, width, height, text, font_size=18,
                    bold=False, color=None, alignment=PP_ALIGN.LEFT, font_name="微软雅黑"):
        """添加文本框"""
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.name = font_name
        p.font.bold = bold
        p.font.color.rgb = color or self.主题色_灰
        p.alignment = alignment
        return tf

    def _创建封面页(self, slide, 页面: 课件页面):
        """创建封面页"""
        self._设置幻灯片背景(slide, self.主题色_青)
        self._添加文本框(slide, 1, 1.5, 11, 1.5, 页面.标题,
                         font_size=40, bold=True, color=self.主题色_白, alignment=PP_ALIGN.CENTER)
        self._添加文本框(slide, 1, 3.5, 11, 1, "基于循证营养学的科普课程",
                         font_size=20, color=RGBColor(0xE0, 0xE8, 0xF0), alignment=PP_ALIGN.CENTER)
        self._添加文本框(slide, 1, 5.5, 11, 0.5, "营养课程生成器 v0.1 · 中文版",
                         font_size=12, color=RGBColor(0xB0, 0xC0, 0xD0), alignment=PP_ALIGN.CENTER)

    def _创建内容页(self, slide, 页面: 课件页面):
        """创建正文内容页"""
        self._设置幻灯片背景(slide, self.主题色_浅灰)
        # 标题栏
        self._添加文本框(slide, 0.5, 0.3, 12, 0.8, 页面.标题,
                         font_size=28, bold=True, color=self.主题色_青)
        # 分隔线
        line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.1), Inches(12), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = self.主题色_青
        line.line.fill.background()
        # 正文
        self._添加文本框(slide, 0.7, 1.3, 11.5, 5.5, 页面.正文.replace('#', '').replace('*', ''),
                         font_size=16, color=self.主题色_灰)
        # 页脚
        self._添加文本框(slide, 0.5, 7, 6, 0.3, f"第{页面.页码}页",
                         font_size=10, color=RGBColor(0x99, 0x99, 0x99))

    def _创建总结页(self, slide, 页面: 课件页面):
        """创建总结/误区/参考文献页"""
        self._设置幻灯片背景(slide, self.主题色_浅灰)
        self._添加文本框(slide, 0.5, 0.3, 12, 0.8, 页面.标题,
                         font_size=28, bold=True, color=self.主题色_金)
        line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.1), Inches(12), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = self.主题色_金
        line.line.fill.background()
        # 简化正文（去除 Markdown 标记）
        简化正文 = 页面.正文.replace('**', '').replace('#', '').replace('*', '')
        self._添加文本框(slide, 0.7, 1.3, 11.5, 5.5, 简化正文,
                         font_size=16, color=self.主题色_灰)
        self._添加文本框(slide, 0.5, 7, 6, 0.3, f"第{页面.页码}页",
                         font_size=10, color=RGBColor(0x99, 0x99, 0x99))

    def 导出(self, 课程: 课程规范, 输出路径: str = "营养课程.pptx"):
        """导出完整的 PPTX 文件"""
        for 页面 in 课程.页面列表:
            slide_layout = self.prs.slide_layouts[6]  # 空白布局
            slide = self.prs.slides.add_slide(slide_layout)

            if 页面.页面类型 == "封面":
                self._创建封面页(slide, 页面)
            elif 页面.页面类型 in ("总结", "误区", "参考文献"):
                self._创建总结页(slide, 页面)
            else:
                self._创建内容页(slide, 页面)

        self.prs.save(输出路径)
        return 输出路径
